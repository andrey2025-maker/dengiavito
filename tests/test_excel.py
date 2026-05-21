from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bot.config import Settings
from bot.services.excel_export import build_excel_backup
from bot.services.excel_scheduler import _parse_report_time


@pytest.mark.asyncio
async def test_build_excel_backup(tmp_path, db):
    settings = Settings(
        bot_token="x",
        owner_id=1,
        db_path=str(tmp_path / "db.sqlite"),
        excel_export_dir=str(tmp_path / "exports"),
    )
    aid = await db.add_apartment("Леонова", "ул. 1")
    await db.update_apartment_door_code(aid, "1234")
    await db.add_booking_range(aid, [date(2026, 5, 10)], 5000, per_day=True)
    await db.add_finance(aid, "expense", date(2026, 5, 1), 1000, "Аренда", None)

    path = await build_excel_backup(db, settings)
    assert path.exists()
    wb = load_workbook(path)
    assert "Квартиры" in wb.sheetnames
    assert "Брони" in wb.sheetnames
    assert "Финансы" in wb.sheetnames
    assert "Админы" in wb.sheetnames
    assert "Сводка" in wb.sheetnames
    ws = wb["Квартиры"]
    assert ws.cell(row=2, column=2).value == "Леонова"


def test_parse_report_time():
    assert _parse_report_time("08:00") == (8, 0)
    assert _parse_report_time("23:59") == (23, 59)
