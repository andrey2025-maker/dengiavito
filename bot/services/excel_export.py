from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from bot.config import Settings, get_excel_zone
from bot.database import Database

TYPE_RU = {"income": "Доход", "expense": "Расход"}


def _sheet_headers(ws, headers: list[str]) -> None:
    bold = Font(bold=True)
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = bold


def _autosize_columns(ws, max_width: int = 42) -> None:
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        length = max(len(str(c.value or "")) for c in column_cells)
        ws.column_dimensions[letter].width = min(max(length + 2, 10), max_width)


async def build_excel_backup(db: Database, settings: Settings) -> Path:
    zone = get_excel_zone(settings)
    now = datetime.now(zone)
    export_dir = Path(settings.excel_export_dir)
    export_dir.mkdir(parents=True, exist_ok=True)
    filename = f"backup_{now.strftime('%Y-%m-%d_%H-%M')}.xlsx"
    path = export_dir / filename

    wb = Workbook()
    wb.remove(wb.active)

    # Квартиры
    ws_apt = wb.create_sheet("Квартиры")
    _sheet_headers(
        ws_apt,
        ["ID", "Название", "Адрес", "Код", "Создано"],
    )
    for i, apt in enumerate(await db.list_apartments(), start=2):
        ws_apt.cell(row=i, column=1, value=apt["id"])
        ws_apt.cell(row=i, column=2, value=apt["name"])
        ws_apt.cell(row=i, column=3, value=apt["address"] or "")
        ws_apt.cell(row=i, column=4, value=apt["door_code"] or "")
        ws_apt.cell(row=i, column=5, value=apt["created_at"])
    _autosize_columns(ws_apt)

    # Брони
    ws_bok = wb.create_sheet("Брони")
    _sheet_headers(
        ws_bok,
        [
            "ID",
            "Квартира",
            "Дата",
            "Сумма за день (₽)",
            "Группа",
            "Создано",
        ],
    )
    for i, row in enumerate(await db.export_all_bookings(), start=2):
        ws_bok.cell(row=i, column=1, value=row["id"])
        ws_bok.cell(row=i, column=2, value=row["apartment_name"])
        ws_bok.cell(row=i, column=3, value=row["booking_date"])
        ws_bok.cell(row=i, column=4, value=row["amount_per_day"])
        ws_bok.cell(row=i, column=5, value=row["group_id"])
        ws_bok.cell(row=i, column=6, value=row["created_at"])
    _autosize_columns(ws_bok)

    # Финансы
    ws_fin = wb.create_sheet("Финансы")
    _sheet_headers(
        ws_fin,
        [
            "ID",
            "Квартира",
            "Тип",
            "Дата",
            "Сумма (₽)",
            "Категория",
            "Комментарий",
            "Создано",
        ],
    )
    for i, row in enumerate(await db.export_all_finances(), start=2):
        ws_fin.cell(row=i, column=1, value=row["id"])
        ws_fin.cell(row=i, column=2, value=row["apartment_name"] or "")
        ws_fin.cell(
            row=i,
            column=3,
            value=TYPE_RU.get(row["record_type"], row["record_type"]),
        )
        ws_fin.cell(row=i, column=4, value=row["record_date"])
        ws_fin.cell(row=i, column=5, value=row["amount"])
        ws_fin.cell(row=i, column=6, value=row["category"] or "")
        ws_fin.cell(row=i, column=7, value=row["comment"] or "")
        ws_fin.cell(row=i, column=8, value=row["created_at"])
    _autosize_columns(ws_fin)

    # Админы
    ws_adm = wb.create_sheet("Админы")
    _sheet_headers(ws_adm, ["User ID", "Имя", "Владелец", "Создано"])
    for i, adm in enumerate(await db.get_admins(), start=2):
        ws_adm.cell(row=i, column=1, value=adm["user_id"])
        ws_adm.cell(row=i, column=2, value=adm["name"])
        ws_adm.cell(row=i, column=3, value="Да" if adm["is_owner"] else "Нет")
        ws_adm.cell(row=i, column=4, value=adm["created_at"])
    _autosize_columns(ws_adm)

    # Сводка
    ws_info = wb.create_sheet("Сводка", 0)
    ws_info["A1"] = "Резервная копия бота «Деньги»"
    ws_info["A1"].font = Font(bold=True, size=14)
    ws_info["A3"] = "Сформировано"
    ws_info["B3"] = now.strftime("%d.%m.%Y %H:%M")
    ws_info["A4"] = "Часовой пояс"
    ws_info["B4"] = str(zone)
    apts = await db.list_apartments()
    bookings = await db.export_all_bookings()
    finances = await db.export_all_finances()
    ws_info["A6"] = "Квартир"
    ws_info["B6"] = len(apts)
    ws_info["A7"] = "Записей броней"
    ws_info["B7"] = len(bookings)
    ws_info["A8"] = "Записей финансов"
    ws_info["B8"] = len(finances)
    _autosize_columns(ws_info, max_width=50)

    wb.save(path)
    return path
